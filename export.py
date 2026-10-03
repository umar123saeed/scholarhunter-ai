"""CSV/Excel export helpers."""
from io import BytesIO
import pandas as pd
from openpyxl.styles import Alignment,Font,PatternFill
from openpyxl.utils import get_column_letter
from tracker import COLUMNS
FILLS={"Critical":"F8B4B4","Soon":"FDE7A8","Upcoming":"BFE8C4","Needs Verification":"D9D9D9","Rolling":"DDEBF7","Expired":"E7E6E6"}
def to_csv_bytes(df):return df[[c for c in COLUMNS if c in df.columns]].to_csv(index=False).encode("utf-8-sig")
def to_excel_bytes(df,gap=None):
    out=BytesIO();cols=[c for c in COLUMNS if c in df.columns];sheet=df[cols].copy()
    if "days_remaining" in sheet:sheet["days_remaining"]=sheet["days_remaining"].astype(object).where(sheet["days_remaining"].notna(),None)
    with pd.ExcelWriter(out,engine="openpyxl") as writer:
        sheet.to_excel(writer,sheet_name="Scholarships",index=False);ws=writer.sheets["Scholarships"]
        for idx,col in enumerate(sheet.columns,1):
            cell=ws.cell(1,idx);cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="1F3A5F");cell.alignment=Alignment(wrap_text=True,vertical="center");ws.column_dimensions[get_column_letter(idx)].width=min(48,max(12,len(col)+4))
        ws.freeze_panes="C2";ws.auto_filter.ref=ws.dimensions
        for r in range(2,ws.max_row+1):
            for c in range(1,ws.max_column+1):ws.cell(r,c).alignment=Alignment(vertical="top",wrap_text=True)
            for name in ("official_link","application_link"):
                if name in sheet.columns:
                    i=list(sheet.columns).index(name)+1;cell=ws.cell(r,i)
                    if cell.value and str(cell.value).startswith("http"):cell.hyperlink=str(cell.value);cell.font=Font(color="0563C1",underline="single")
            if "urgency" in sheet.columns:
                i=list(sheet.columns).index("urgency")+1;cell=ws.cell(r,i)
                if cell.value in FILLS:cell.fill=PatternFill("solid",fgColor=FILLS[cell.value])
        rows=[]
        if gap:rows += [("Strength",x) for x in gap.strengths]+[("Gap",x) for x in gap.gaps]+[("Recommendation",x) for x in gap.recommendations]
        pd.DataFrame(rows,columns=["Type","Detail"]).to_excel(writer,sheet_name="Gap Report",index=False)
    return out.getvalue()
